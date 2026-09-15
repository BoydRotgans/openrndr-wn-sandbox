#!/usr/bin/env python3
"""Turn the deck into a practice sheet: every slide as a small picture with its
speaker note beside it, as a PDF you can hold while you rehearse.

    .venv/bin/python tools/speaker_notes_pdf.py

The source is export/wn-speaker-notes.pptx — the deck tools/figma_to_pptx.py
builds, whose notes are PowerPoint's own notes pane rather than captions drawn on
the slides. Nothing here reads Figma or the network: the pptx already carries both
halves, so this is a pure re-typesetting of what is in it.

The thumbnail is taken from the slide's own picture rather than from the
export/pptx-images cache, so a slide and the icon beside it cannot come apart.
The cache is used only to recover each slide's Figma frame name (`2-07`), by
matching bytes — a name is worth having on a practice sheet, because it is what
you say to whoever is driving the deck.

    --only-notes   leave out the slides that carry no note
    --out PATH     write somewhere other than export/wn-speaker-notes.pdf
"""
import argparse
import base64
import hashlib
import html
import io
import os
import re
import subprocess
import sys
import tempfile
from datetime import date

from PIL import Image
from pptx import Presentation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CHROME = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"

# The chapter cards carry their titles as *drawing* rather than as text — they are
# set in Figma and rasterised — so the four names are stated here. They are the
# show's own, off Slideshow.kt, with chapter 2's soft hyphen closed up.
CHAPTERS = {
    "1": "De wereld van bouwen",
    "2": "Waardekader en verantwoordelijkheid",
    "3": "Beton: ruggengraat en transitie",
    "4": "The circle: een nieuwe manier van denken",
}

CARD_RE = re.compile(r"^chapter-(\d+)$")
FRAME_RE = re.compile(r"^(\d+)-(\d+)$")


# ---------------------------------------------------------------- reading


def frame_names(prs, cache_dir):
    """Each slide's Figma frame name, recovered by matching its picture's bytes
    against the render cache. Returns None for a slide the cache does not hold,
    so a missing or stale cache costs the names and nothing else."""
    if not os.path.isdir(cache_dir):
        return [None] * len(prs.slides)
    by_hash = {}
    for name in os.listdir(cache_dir):
        if not name.endswith(".png"):
            continue
        with open(os.path.join(cache_dir, name), "rb") as f:
            by_hash[hashlib.md5(f.read()).hexdigest()] = re.sub(r"@\d+x\.png$", "", name)
    out = []
    for slide in prs.slides:
        picture = next((s for s in slide.shapes if s.shape_type == 13), None)
        out.append(by_hash.get(hashlib.md5(picture.image.blob).hexdigest()) if picture else None)
    return out


def read_deck(path, cache_dir, thumb_width, quality):
    """The deck as a list of rows: picture, note, and where the slide sits."""
    prs = Presentation(path)
    names = frame_names(prs, cache_dir)
    rows = []
    for index, slide in enumerate(prs.slides):
        picture = next((s for s in slide.shapes if s.shape_type == 13), None)
        note = slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else ""
        rows.append({
            "number": index + 1,
            "name": names[index],
            "note": note,
            "thumb": thumbnail(picture.image.blob, thumb_width, quality) if picture else None,
        })
    return rows


def thumbnail(blob, width, quality):
    """A slide's picture, small enough to carry 78 of them in one file. The deck is
    3840x2160 for a single pane and 7816x2160 for the whole wall, so the height is
    left to follow the width — which also tells the two kinds apart at a glance."""
    image = Image.open(io.BytesIO(blob)).convert("RGB")
    image.thumbnail((width, width), Image.LANCZOS)
    buffer = io.BytesIO()
    image.save(buffer, "JPEG", quality=quality, optimize=True)
    return "data:image/jpeg;base64," + base64.b64encode(buffer.getvalue()).decode()


# ---------------------------------------------------------------- typesetting


# PowerPoint keeps a symbol font's glyph as a private-use code point rather than as
# the character it draws, so a Wingdings arrow arrives as U+F0E0 and prints as a
# hollow box in any face that is not Wingdings. One occurrence in this deck, and it
# is in the middle of a sentence ("Overstap CEM I -> CEM II"), so it is translated
# rather than dropped.
SYMBOLS = {"\uf0e0": "\u2192", "\uf0e8": "\u2192", "\uf0a7": "\u2022", "\uf0b7": "\u2022"}


def paragraphs(note):
    """A note as it was written: blank lines part it, single ones break within it."""
    for glyph, char in SYMBOLS.items():
        note = note.replace(glyph, char)
    blocks = [b for b in re.split(r"\n\s*\n", note) if b.strip()]
    return ["<br>".join(html.escape(line) for line in b.strip().split("\n")) for b in blocks]


def build_html(rows, source, only_notes):
    parts = []
    with_notes = sum(1 for r in rows if r["note"])
    parts.append(HEAD)
    parts.append(
        f"<header class='sheet'><h1>Speaker notes</h1>"
        f"<p class='meta'>{html.escape(os.path.basename(source))} &middot; "
        f"{len(rows)} slides &middot; {with_notes} with notes &middot; "
        f"{date.today().isoformat()}</p></header>"
    )

    chapter = None
    # The deck repeats a note across the build steps of one picture — six of these
    # slides carry the note of the one before them, word for word. Printed again it
    # reads as a new thing to say, so a repeat is marked instead of restated, and
    # the comparison skips the silent slides between: a note that comes back after
    # a build step with nothing on it is still a note you have already read.
    spoken, spoken_at = None, None
    for row in rows:
        card = CARD_RE.match(row["name"] or "")
        if card:
            # A chapter card is a divider rather than a row: it is the same wall the
            # audience sees, and what is said over it is the chapter's name.
            chapter = card.group(1)
            parts.append(
                f"<section class='chapter'>"
                f"<div class='chapter-no'>{html.escape(chapter)}</div>"
                f"<div class='chapter-title'>{html.escape(CHAPTERS.get(chapter, ''))}</div>"
                f"<img class='chapter-thumb' src='{row['thumb']}' alt=''>"
                f"</section>"
            )
            continue

        if only_notes and not row["note"]:
            continue

        note = row["note"]
        label = row["name"] or f"slide {row['number']}"
        repeat = bool(note) and note == spoken
        if note and not repeat:
            body = "".join(f"<p>{p}</p>" for p in paragraphs(note))
        elif repeat:
            body = f"<p class='again'>&uarr; zelfde notitie als {html.escape(spoken_at)}</p>"
        else:
            body = "<p class='silent'>&mdash;</p>"
        if note:
            spoken, spoken_at = note, label

        parts.append(
            f"<article class='row{'' if note and not repeat else ' quiet'}'>"
            f"<div class='pic'><img src='{row['thumb']}' alt=''>"
            f"<div class='tag'><span class='frame'>{html.escape(label)}</span>"
            f"<span class='num'>{row['number']}</span></div></div>"
            f"<div class='note'>{body}</div>"
            f"</article>"
        )

    parts.append("</body></html>")
    return "".join(parts)


HEAD = """<!doctype html><meta charset="utf-8"><title>Speaker notes</title><style>
@page { size: A4 portrait; margin: 14mm 13mm 16mm; }
* { box-sizing: border-box; }
body { margin: 0; font: 10pt/1.45 "Helvetica Neue", Helvetica, Arial, sans-serif; color: #111; }
h1 { font-family: Rockwell, Georgia, serif; font-size: 22pt; margin: 0 0 2mm; letter-spacing: .01em; }
.meta { margin: 0; font-size: 8.5pt; color: #777; }
header.sheet { padding-bottom: 4mm; border-bottom: 1.5pt solid #111; margin-bottom: 6mm; }

/* A chapter card is the wall the audience sees, so it stands across the sheet
   rather than in the picture column, and starts a page of its own. */
section.chapter { break-before: page; break-inside: avoid; margin: 0 0 7mm;
  padding-bottom: 4mm; border-bottom: 1.5pt solid #111; }
section.chapter:first-of-type { break-before: auto; }
.chapter-no { font: 8.5pt/1 "Helvetica Neue", Arial, sans-serif; letter-spacing: .16em;
  text-transform: uppercase; color: #999; margin-bottom: 1.5mm; }
.chapter-title { font-family: Rockwell, Georgia, serif; font-size: 19pt; line-height: 1.15;
  margin-bottom: 3.5mm; }
.chapter-thumb { width: 100%; display: block; }

article.row { display: flex; gap: 6mm; break-inside: avoid; page-break-inside: avoid;
  padding: 3.5mm 0; border-bottom: .4pt solid #ddd; align-items: flex-start; }
article.row.quiet { padding: 2mm 0; }
.pic { flex: 0 0 62mm; }
.pic img { width: 100%; display: block; border: .4pt solid #ccc; }
.tag { display: flex; justify-content: space-between; margin-top: 1.2mm;
  font-size: 7.5pt; color: #888; letter-spacing: .06em; }
.frame { font-weight: 600; color: #444; }
.note { flex: 1 1 auto; min-width: 0; }
.note p { margin: 0 0 2.4mm; }
.note p:last-child { margin-bottom: 0; }
.note .silent { color: #ccc; }
.note .again { color: #999; font-size: 8.5pt; font-style: italic; }
article.row.quiet .pic { flex: 0 0 42mm; }
</style><body>"""


# ---------------------------------------------------------------- output


def to_pdf(markup, out):
    """Chrome prints it. It is the one renderer on this machine that takes the
    print CSS above — the page size, the margins, and `break-inside: avoid`, which
    is what keeps a note and its picture on one page."""
    with tempfile.TemporaryDirectory() as tmp:
        page = os.path.join(tmp, "notes.html")
        with open(page, "w") as f:
            f.write(markup)
        subprocess.run([
            CHROME, "--headless", "--disable-gpu", "--no-pdf-header-footer",
            "--virtual-time-budget=20000",
            f"--print-to-pdf={out}", f"file://{page}",
        ], check=True, capture_output=True)


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--pptx", default=os.path.join(ROOT, "export/wn-speaker-notes.pptx"))
    ap.add_argument("--cache", default=os.path.join(ROOT, "export/pptx-images"))
    ap.add_argument("--out", default=os.path.join(ROOT, "export/wn-speaker-notes.pdf"))
    ap.add_argument("--only-notes", action="store_true",
                    help="leave out the slides that carry no note")
    ap.add_argument("--width", type=int, default=760, help="thumbnail width in pixels")
    ap.add_argument("--quality", type=int, default=78)
    args = ap.parse_args()

    if not os.path.isfile(args.pptx):
        sys.exit(f"No deck at {args.pptx} — run tools/figma_to_pptx.py first.")
    if not os.path.isfile(CHROME):
        sys.exit(f"No Chrome at {CHROME}, and it is what prints the sheet.")

    rows = read_deck(args.pptx, args.cache, args.width, args.quality)
    markup = build_html(rows, args.pptx, args.only_notes)
    os.makedirs(os.path.dirname(args.out), exist_ok=True)
    to_pdf(markup, args.out)

    shown = sum(1 for r in rows if not CARD_RE.match(r["name"] or "")
                and (r["note"] or not args.only_notes))
    print(f"-> {os.path.relpath(args.out, ROOT)}: {shown} slides, "
          f"{sum(1 for r in rows if r['note'])} with notes, "
          f"{os.path.getsize(args.out) // 1024} KB")


if __name__ == "__main__":
    main()
